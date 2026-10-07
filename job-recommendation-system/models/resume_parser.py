"""Bounded PDF/DOCX text extraction with an explicit user-review step."""

import io
import re
import zipfile
from pathlib import Path
from docx import Document
from pypdf import PdfReader
from werkzeug.utils import secure_filename
from models.skill_extractor import extract_skills

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_TEXT_CHARS = 50000


class ResumeError(ValueError):
    pass


def parse_resume(filename, data):
    safe_name = secure_filename(filename or "")
    suffix = Path(safe_name).suffix.lower()
    if suffix not in {".pdf", ".docx"}:
        raise ResumeError("Upload a PDF or DOCX resume.")
    if not data or len(data) > MAX_FILE_BYTES:
        raise ResumeError("The resume must contain data and be no larger than 5 MB.")
    try:
        if suffix == ".pdf":
            if not data.startswith(b"%PDF-"):
                raise ResumeError("This file is not a valid PDF.")
            reader = PdfReader(io.BytesIO(data), strict=False)
            if reader.is_encrypted:
                raise ResumeError("Please upload an unencrypted PDF.")
            if len(reader.pages) > 30:
                raise ResumeError("A resume may contain at most 30 pages.")
            chunks = []
            for page in reader.pages:
                if (
                    len(page.get_contents().get_data() if page.get_contents() else b"")
                    > 5 * 1024 * 1024
                ):
                    raise ResumeError(
                        "This PDF page is too complex. Export a simpler text PDF."
                    )
                chunks.append(page.extract_text() or "")
                if sum(map(len, chunks)) > MAX_TEXT_CHARS:
                    break
            text = "\n".join(chunks)
        else:
            if not zipfile.is_zipfile(io.BytesIO(data)):
                raise ResumeError("This file is not a valid DOCX document.")
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                entries = archive.infolist()
                if (
                    len(entries) > 300
                    or sum(entry.file_size for entry in entries) > 20 * 1024 * 1024
                ):
                    raise ResumeError(
                        "The expanded DOCX document exceeds the safe size limit."
                    )
                if "word/document.xml" not in archive.namelist():
                    raise ResumeError("The DOCX file has no document content.")
            document = Document(io.BytesIO(data))
            text = "\n".join(
                [p.text for p in document.paragraphs]
                + [
                    " ".join(cell.text for cell in row.cells)
                    for table in document.tables
                    for row in table.rows
                ]
            )
    except ResumeError:
        raise
    except Exception as exc:
        raise ResumeError(
            "The document could not be read. Re-export it as PDF or DOCX."
        ) from exc
    text = text[:MAX_TEXT_CHARS].strip()
    if len(text) < 30:
        raise ResumeError(
            "No readable text found. Scanned/image resumes need OCR before upload."
        )
    return {"filename": safe_name[:200], **extract_resume_fields(text)}


def section_lines(text, headings):
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    result, capturing = [], False
    all_headings = {
        "skills",
        "education",
        "experience",
        "work experience",
        "projects",
        "certifications",
        "certificates",
        "summary",
        "contact",
        "interests",
    }
    for line in lines:
        key = line.casefold().strip(": ")
        if key in all_headings:
            capturing = key in headings
        elif capturing:
            result.append(line)
    return "\n".join(result)[:5000]


def extract_resume_fields(text):
    degrees = re.findall(
        r"\b(?:B\.?\s?Tech|B\.?\s?Sc|BCA|B\.?\s?E\.?|B\.?\s?A\.?|MCA|M\.?\s?Tech|MBA|M\.?\s?Sc|M\.?\s?A\.?|Diploma|Bachelor[^\n,.]{0,65}|Master[^\n,.]{0,65}|12th|10th)\b[^\n]{0,65}",
        text,
        re.I,
    )
    years = re.findall(
        r"\b(\d{1,2}(?:\.\d)?)\+?\s*(?:years?|yrs?)\s+(?:of\s+)?(?:professional\s+|work\s+|industry\s+)?experience\b",
        text,
        re.I,
    )
    roles = re.findall(
        r"\b(?:Data Analyst|Business Analyst|Data Scientist|Python Developer|Software Engineer|Power BI Developer|Customer Support Executive|Graphic Designer|Cloud Engineer|Cybersecurity Analyst)\b",
        text,
        re.I,
    )
    warnings = [
        "Review every extracted field. Dates and job durations are not automatically summed."
    ]
    if not years:
        warnings.append(
            "Total experience was not stated explicitly; please enter it yourself."
        )
    return {
        "text": text,
        "skills": extract_skills(text),
        "education": degrees[0] if degrees else "",
        "experience_years": min(float(max(years, key=float)), 60) if years else None,
        "certifications": section_lines(text, {"certifications", "certificates"}),
        "projects": section_lines(text, {"projects"}),
        "experience_summary": section_lines(text, {"experience", "work experience"}),
        "preferred_role": roles[0].title() if roles else "",
        "job_titles": sorted(set(roles)),
        "warnings": warnings,
    }
